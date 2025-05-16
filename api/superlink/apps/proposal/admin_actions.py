from datetime import datetime

from openpyxl.utils import get_column_letter
import openpyxl
from openpyxl.styles import Font, Alignment
from django.contrib.admin import ModelAdmin
from django.http import HttpResponse
from openpyxl import Workbook
from rest_framework import status
from rest_framework.request import Request
from apps.proposal.models import Proposal


def export_to_xls(model_admin: ModelAdmin, request: Request, query_set: Proposal) -> HttpResponse:
    """Export the selected proposals to an XLSX file"""
    model_admin.message_user(request, "Заявки выгружены успешно.")

    try:
        # Define the fields to include and headers to display
        display_headers: dict[str, str] = {
            'created_at': 'Дата создания',
            'get_product': 'Продукт',
            'source': 'сточник',
            'get_applicant_fullname': 'ФИО',
            'applicant_phone_number': 'Номер телефона',
            'status': 'Статус',
            'get_user_log_entry': 'Обработано',
            'get_user_log_entry_datetime': 'Дата и время обработки',
        }

        fields_to_include: list[str] = list(display_headers.keys())

        # Create a new Excel workbook
        wb: Workbook = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Заявки"

        title_font: Font = Font(size=16, bold=True)
        description_font: Font = Font(size=12)
        metadata_font: Font = Font(size=10, italic=True)

        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = "Отчет по заявкам (VTiger)"
        title_cell.font = title_font
        title_cell.alignment = Alignment(horizontal='center')

        ws.merge_cells('A2:E2')
        description_cell = ws['A2']
        description_cell.value = "Этот отчет содержит заявки на консультации."
        description_cell.font = description_font
        description_cell.alignment = Alignment(horizontal='center')

        # Metadata for creation time
        ws['A3'] = f"Создано: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        ws['A3'].font = metadata_font

        # Headers Row
        header_font: Font = Font(bold=True)
        for col_num, field in enumerate(fields_to_include, 1):
            header_value = display_headers.get(field, field)
            cell = ws.cell(row=6, column=col_num, value=header_value)
            cell.font = header_font

        # Populate data rows for the selected query_set
        for row_num, proposal in enumerate(query_set, 7):
            for col_num, field in enumerate(fields_to_include, 1):
                if field == 'status':
                    value = proposal.get_status_display()
                else:
                    value = getattr(proposal, field, '-')
                ws.cell(row=row_num, column=col_num, value=value)

        # Adjust column widths
        for col_num in range(1, len(fields_to_include) + 1):
            column_letter = get_column_letter(col_num)
            max_length = 0

            for row in ws.iter_rows(min_row=6, max_row=ws.max_row, min_col=col_num, max_col=col_num):
                for cell in row:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception as exc:
                        print(str(exc))
                        pass

            adjusted_width = (max_length + 2) * 1.2
            ws.column_dimensions[column_letter].width = adjusted_width

        # Prepare response to download the file
        file_name: str = f'bakai-kg-proposals__{datetime.now():%Y-%m-%d_%H_%M_%S}.xlsx'
        response: HttpResponse = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename={file_name}'
        wb.save(response)

        return response

    except Exception as exc:
        model_admin.message_user(request, f"Внутреняя ошибка: {str(exc)}", level="error")
        return HttpResponse(status=status.HTTP_500_INTERNAL_SERVER_ERROR)
